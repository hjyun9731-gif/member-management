"""화물운송종사자격증 번호(0-00-000000) — 회원/예정자/양도양수/도내양도 저장·조회·전환·엑셀 검증.
가상 데이터 + 임시 sqlite 사용 (운영 DB 접속 없음)."""
import io, os, sys
os.environ.setdefault("DATABASE_URL", "sqlite:////tmp/mm_fqn_test.db")
sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
import pytest
from fastapi import FastAPI
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from app import models
from app.auth import get_current_user, require_admin
from app.database import Base, engine, SessionLocal, get_db
from app.routers import candidates, members, transfer_ledger
from app.services.freight_qualification import normalize_fqn


@pytest.fixture()
def client():
    Base.metadata.drop_all(bind=engine); Base.metadata.create_all(bind=engine)
    app = FastAPI()
    app.include_router(candidates.router, prefix="/api/candidates")
    app.include_router(members.router, prefix="/api/members")
    app.include_router(transfer_ledger.router, prefix="/api/transfer-ledger")

    def _db():
        s = SessionLocal()
        try:
            yield s
        finally:
            s.close()
    user = models.User(id=1, username="t", password_hash="x", role="admin", full_name="t")
    app.dependency_overrides[get_db] = _db
    app.dependency_overrides[get_current_user] = lambda: user
    app.dependency_overrides[require_admin] = lambda: user
    return TestClient(app, raise_server_exceptions=False)


def _db():
    return SessionLocal()


# ───────── 형식 ─────────
@pytest.mark.parametrize("raw,exp", [
    ("1-23-456789", "1-23-456789"), ("123456789", "1-23-456789"), ("012345678", "0-12-345678"),
    (" 0 12 345678 ", "0-12-345678"), ("", ""), (None, ""), ("   ", ""),
])
def test_normalize_ok(raw, exp):
    assert normalize_fqn(raw) == exp


@pytest.mark.parametrize("raw", ["1234", "12345678", "1234567890", "1-23-4567890"])
def test_normalize_rejects_non_9_digits(raw):
    with pytest.raises(ValueError):
        normalize_fqn(raw)


# ───────── 회원 ─────────
def test_member_create_update_detail_keep_leading_zero(client):
    r = client.post("/api/members", json={"vehicle_number": "강원80배1001", "name": "홍길동", "freight_qualification_number": "012345678"})
    assert r.status_code == 200, r.text
    mid = r.json()["id"]
    assert r.json()["freight_qualification_number"] == "0-12-345678"
    assert client.get(f"/api/members/{mid}").json()["freight_qualification_number"] == "0-12-345678"
    # 기존 값 수정
    r = client.put(f"/api/members/{mid}", json={"freight_qualification_number": "9-87-654321"})
    assert r.status_code == 200 and r.json()["freight_qualification_number"] == "9-87-654321"
    # 다른 필드만 수정하면 번호는 그대로
    client.put(f"/api/members/{mid}", json={"memo": "x"})
    assert client.get(f"/api/members/{mid}").json()["freight_qualification_number"] == "9-87-654321"
    # 잘못된 형식은 거부되고 기존 값 유지
    assert client.put(f"/api/members/{mid}", json={"freight_qualification_number": "123"}).status_code == 400
    assert client.get(f"/api/members/{mid}").json()["freight_qualification_number"] == "9-87-654321"
    # 선택 항목: 비우기 가능
    client.put(f"/api/members/{mid}", json={"freight_qualification_number": ""})
    assert client.get(f"/api/members/{mid}").json()["freight_qualification_number"] == ""


def test_member_created_without_number_can_add_later(client):
    mid = client.post("/api/members", json={"vehicle_number": "강원80배1002", "name": "기존회원"}).json()["id"]
    assert client.get(f"/api/members/{mid}").json()["freight_qualification_number"] == ""
    client.put(f"/api/members/{mid}", json={"freight_qualification_number": "1-23-456789"})
    assert client.get(f"/api/members/{mid}").json()["freight_qualification_number"] == "1-23-456789"


def test_member_excel_export_text_keeps_leading_zero(client):
    client.post("/api/members", json={"vehicle_number": "강원80배1003", "name": "엑셀", "freight_qualification_number": "0-12-345678"})
    r = client.get("/api/members/export/excel")
    assert r.status_code == 200
    ws = load_workbook(io.BytesIO(r.content)).active
    hdr = [c.value for c in ws[1]]
    col = hdr.index("freight_qualification_number")
    cell = ws.cell(row=2, column=col + 1)
    assert cell.value == "0-12-345678" and cell.number_format == "@"


# ───────── 예정자 ─────────
def test_candidate_save_edit_and_register_keeps_number(client):
    r = client.post("/api/candidates", json={"vehicle_number": "강원80배2001", "name": "예정자", "freight_qualification_number": "023456789"})
    assert r.status_code == 200, r.text
    cid = r.json()["id"]
    assert r.json()["freight_qualification_number"] == "0-23-456789"
    assert client.get(f"/api/candidates/{cid}").json()["freight_qualification_number"] == "0-23-456789"
    assert client.put(f"/api/candidates/{cid}", json={"freight_qualification_number": "5-55-555555"}).json()["freight_qualification_number"] == "5-55-555555"
    assert client.put(f"/api/candidates/{cid}", json={"freight_qualification_number": "12"}).status_code == 400
    r = client.post(f"/api/candidates/{cid}/register", json={"approval_date": "2026-10-01", "management_number": "신26-9001"})
    assert r.status_code == 200, r.text
    mid = r.json()["member_id"]
    assert client.get(f"/api/members/{mid}").json()["freight_qualification_number"] == "5-55-555555"  # 전환 후 유지


# ───────── 양도양수대장 ─────────
def test_transfer_ledger_create_update_view_export(client):
    r = client.post("/api/transfer-ledger", json={"vehicle_number": "강원80배3001", "transferor": "갑", "transferee": "을", "freight_qualification_number": "0-12-345678"})
    assert r.status_code == 200, r.text
    tid = r.json()["id"]
    assert client.get(f"/api/transfer-ledger/{tid}").json()["freight_qualification_number"] == "0-12-345678"
    assert client.put(f"/api/transfer-ledger/{tid}", json={"freight_qualification_number": "3-33-333333"}).json()["freight_qualification_number"] == "3-33-333333"
    assert client.put(f"/api/transfer-ledger/{tid}", json={"freight_qualification_number": "99999999999"}).status_code == 400
    ws = load_workbook(io.BytesIO(client.get("/api/transfer-ledger/export/excel").content)).active
    hdr = [c.value for c in ws[1]]
    assert "freight_qualification_number" in hdr
    assert ws.cell(row=2, column=hdr.index("freight_qualification_number") + 1).value == "3-33-333333"


def test_ledger_register_member_copies_number_to_transferee_member(client):
    tid = client.post("/api/transfer-ledger", json={"vehicle_number": "강원80배3002", "transferor": "갑", "transferee": "을", "approval_date": "2026-10-01", "freight_qualification_number": "012345678"}).json()["id"]
    r = client.post(f"/api/transfer-ledger/{tid}/register-member", json={"management_number": "양26-9002"})
    assert r.status_code == 200, r.text
    mid = r.json()["member_id"] if "member_id" in r.json() else None
    db = _db()
    try:
        m = db.query(models.LicenseHolder).filter(models.LicenseHolder.management_number == "양26-9002").first()
        assert m is not None and m.freight_qualification_number == "0-12-345678"
    finally:
        db.close()


def test_ledger_edit_syncs_to_transferee_but_never_blanks_it(client):
    db = _db()
    try:
        m = models.LicenseHolder(vehicle_number="강원80배3003", name="양수자", status="active", freight_qualification_number="4-44-444444")
        db.add(m); db.commit(); db.refresh(m)
        t = models.TransferLedger(vehicle_number="강원80배3003", transferor="양도자", transferee="양수자", transferee_member_id=m.id, member_id=m.id)
        db.add(t); db.commit(); db.refresh(t); mid, tid = m.id, t.id
    finally:
        db.close()
    # 대장 폼이 빈 값이면 회원 번호를 지우지 않는다
    assert client.put(f"/api/transfer-ledger/{tid}", json={"transferee": "양수자", "freight_qualification_number": ""}).status_code == 200
    assert client.get(f"/api/members/{mid}").json()["freight_qualification_number"] == "4-44-444444"
    # 대장에서 새 번호를 넣으면 양수자 회원에 반영
    assert client.put(f"/api/transfer-ledger/{tid}", json={"freight_qualification_number": "6-66-666666"}).status_code == 200
    assert client.get(f"/api/members/{mid}").json()["freight_qualification_number"] == "6-66-666666"


# ───────── 도내 양도양수: 양도인/양수인 번호 분리 ─────────
def _transferor(num="1-11-111111", plate="강원80배4001"):
    db = _db()
    try:
        m = models.LicenseHolder(vehicle_number=plate, name="양도인", category="택배", status="active",
                                 management_number="신20-1", freight_qualification_number=num)
        db.add(m); db.commit(); db.refresh(m); return m.id
    finally:
        db.close()


def _body(tid, **kw):
    b = {"transferor_member_id": tid, "closure_date": "2026-10-20", "transferee_target": "member",
         "transferee_name": "양수인", "transferee_vehicle_number": "강원80배4001", "transferee_region": "춘천"}
    b.update(kw); return b


def test_domestic_transfer_keeps_transferor_and_transferee_numbers_separate(client):
    tid = _transferor("1-11-111111")
    r = client.post("/api/transfer-ledger/domestic", json=_body(tid, transferee_freight_qualification_number="022222222"))
    assert r.status_code == 200, r.text
    res = r.json()
    db = _db()
    try:
        transferor = db.get(models.LicenseHolder, tid)
        transferee = db.get(models.LicenseHolder, res["transferee_member_id"])
        ledger = db.get(models.TransferLedger, res["transfer_ledger_id"])
        assert transferor.status == "closed" and transferor.freight_qualification_number == "1-11-111111"
        assert transferee.freight_qualification_number == "0-22-222222"
        assert ledger.freight_qualification_number == "0-22-222222"          # 대장 = 양수인 번호
        assert ledger.transferor_member_id == tid and ledger.transferee_member_id == transferee.id
    finally:
        db.close()


def test_domestic_transfer_without_transferee_number_never_copies_transferor_number(client):
    tid = _transferor("1-11-111111")
    r = client.post("/api/transfer-ledger/domestic", json=_body(tid))
    assert r.status_code == 200, r.text
    db = _db()
    try:
        transferee = db.get(models.LicenseHolder, r.json()["transferee_member_id"])
        ledger = db.get(models.TransferLedger, r.json()["transfer_ledger_id"])
        assert (transferee.freight_qualification_number or "") == ""
        assert (ledger.freight_qualification_number or "") == ""
    finally:
        db.close()


def test_domestic_transfer_candidate_target_stores_number_on_candidate(client):
    tid = _transferor("1-11-111111")
    r = client.post("/api/transfer-ledger/domestic", json=_body(tid, transferee_target="candidate", transferee_freight_qualification_number="7-77-777777"))
    assert r.status_code == 200, r.text
    db = _db()
    try:
        cand = db.get(models.Candidate, r.json()["transferee_candidate_id"])
        assert cand.freight_qualification_number == "7-77-777777"
        assert db.get(models.LicenseHolder, tid).freight_qualification_number == "1-11-111111"
    finally:
        db.close()


def test_domestic_transfer_invalid_number_rejected_and_nothing_saved(client):
    tid = _transferor()
    r = client.post("/api/transfer-ledger/domestic", json=_body(tid, transferee_freight_qualification_number="12345"))
    assert r.status_code == 400
    db = _db()
    try:
        assert db.get(models.LicenseHolder, tid).status == "active"
        assert db.query(models.TransferLedger).count() == 0
        assert db.query(models.Closure).count() == 0
        assert db.query(models.LicenseHolder).count() == 1
    finally:
        db.close()


def test_domestic_transfer_link_existing_conflicting_number_rolls_back_everything(client):
    tid = _transferor("1-11-111111")
    db = _db()
    try:
        ex = models.LicenseHolder(vehicle_number="강원80배4999", name="기존양수인", status="active", management_number="신25-5",
                                  freight_qualification_number="8-88-888888")
        db.add(ex); db.commit(); db.refresh(ex); exid = ex.id
    finally:
        db.close()
    r = client.post("/api/transfer-ledger/domestic", json=_body(tid, link_existing_id=exid, link_existing_type="member",
                                                               transferee_freight_qualification_number="9-99-999999"))
    assert r.status_code == 400, r.text
    db = _db()
    try:
        assert db.get(models.LicenseHolder, tid).status == "active"          # 양도자 폐업 안 됨
        assert db.query(models.TransferLedger).count() == 0                   # 대장 없음
        assert db.query(models.Closure).count() == 0                          # 폐업현황 없음
        assert db.get(models.LicenseHolder, exid).freight_qualification_number == "8-88-888888"
    finally:
        db.close()


def test_domestic_transfer_link_existing_fills_missing_number_only(client):
    tid = _transferor("1-11-111111")
    db = _db()
    try:
        ex = models.LicenseHolder(vehicle_number="강원80배4998", name="기존양수인", status="active", management_number="신25-6")
        db.add(ex); db.commit(); db.refresh(ex); exid = ex.id
    finally:
        db.close()
    r = client.post("/api/transfer-ledger/domestic", json=_body(tid, link_existing_id=exid, link_existing_type="member",
                                                               transferee_freight_qualification_number="2-22-222222"))
    assert r.status_code == 200, r.text
    db = _db()
    try:
        assert db.get(models.LicenseHolder, exid).freight_qualification_number == "2-22-222222"
        assert db.get(models.TransferLedger, r.json()["transfer_ledger_id"]).freight_qualification_number == "2-22-222222"
        assert db.get(models.LicenseHolder, tid).freight_qualification_number == "1-11-111111"
    finally:
        db.close()
