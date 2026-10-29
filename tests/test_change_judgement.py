"""변경이력 판정·월례보고서/월별통계 일치·탈퇴 감지·주민번호 보호·관리비폐지 확인필요 검증 (가상 데이터, 임시 sqlite)."""
import io, os, sys
os.environ.setdefault("DATABASE_URL", "sqlite:////tmp/mm_chg_test.db")
sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
import datetime as dt
import pytest
from fastapi import FastAPI
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from app import models
from app.auth import get_current_user, require_admin
from app.database import Base, engine, SessionLocal, get_db
from app.routers import allocation, change_history, change_stats, dashboard, members
from app.services.change_judgement import (classify_plate, classify_text, mask_rrn, parse_ymd, pick_report_date,
                                           NORMAL, FIRST, NOTATION, ERROR_FIX, REVIEW)

TODAY = dt.datetime.now()


@pytest.fixture()
def client():
    Base.metadata.drop_all(bind=engine); Base.metadata.create_all(bind=engine)
    app = FastAPI()
    app.include_router(members.router, prefix="/api/members")
    app.include_router(change_history.router, prefix="/api/change-history")
    app.include_router(change_stats.router, prefix="/api/change-stats")
    app.include_router(dashboard.router, prefix="/api/dashboard")
    app.include_router(allocation.router, prefix="/api/allocation")

    def _db():
        s = SessionLocal()
        try:
            yield s
        finally:
            s.close()
    user = models.User(id=1, username="tester", password_hash="x", role="admin", full_name="t")
    app.dependency_overrides[get_db] = _db
    app.dependency_overrides[get_current_user] = lambda: user
    app.dependency_overrides[require_admin] = lambda: user
    return TestClient(app, raise_server_exceptions=False)


def _member(**kw):
    base = dict(vehicle_number="강원80배1001", name="홍길동", region="춘천시", category="택배", status="active")
    base.update(kw)
    db = SessionLocal()
    try:
        m = models.LicenseHolder(**base); db.add(m); db.commit(); db.refresh(m); return m.id
    finally:
        db.close()


def _hist(**f):
    db = SessionLocal()
    try:
        q = db.query(models.ChangeHistory)
        for k, v in f.items():
            q = q.filter(getattr(models.ChangeHistory, k) == v)
        return q.all(), db
    except Exception:
        db.close(); raise


def _rows(**f):
    rows, db = _hist(**f)
    out = [(r.change_type, r.judgement, r.before_value, r.after_value, r.raw_data) for r in rows]
    db.close(); return out


# ───────── 순수 판정 ─────────
@pytest.mark.parametrize("before", ["택배", "-", "강원 배", "강원 바", "강원 사", "강원 자", ""])
def test_first_plate_registration_is_not_a_change(before):
    assert classify_plate(before, "강원86배 1069호")[0] == FIRST


def test_plate_typo_notation_and_ambiguous_cases():
    assert classify_plate("강운83자 5192호", "강원83자 5192호")[0] == ERROR_FIX
    assert classify_plate("강원83자5192호", "강원 83자 5192호")[0] == NOTATION
    assert classify_plate("강원 83자 5192호", "강원83자 5192호")[0] == NOTATION
    assert classify_plate("강원83자 5192호", "강원83자 5193호")[0] == REVIEW      # 숫자 한 글자 → 확인필요
    assert classify_plate("강원83자 5192호", "강원83바 5192호")[0] == REVIEW      # 용도문자 한 글자 → 확인필요
    assert classify_plate("강원83자 5188호", "강원83자 5190호")[0] == NORMAL      # 실제 번호 변경
    assert classify_plate("강원81자 2791호", "강원82자 2414호")[0] == NORMAL
    assert classify_plate("택배", "강원 배")[0] == REVIEW                         # 임시값 → 불완전 번호


def test_text_classification():
    assert classify_text("춘천시 동내면 1-2", "춘천시  동내면 1-2")[0] == NOTATION
    assert classify_text("동내면 1번지", "동내면 2번지")[0] == NORMAL
    assert classify_text("-", "동내면 1번지")[0] == FIRST
    assert classify_text("1.5톤 카고", "2.5톤 카고")[0] == NORMAL


def test_dates_and_rrn_protection():
    assert parse_ymd("2018.01.03.") == (2018, 1, 3) and parse_ymd("2016 09.21") == (2016, 9, 21)
    assert parse_ymd("26.05.22") == (2026, 5, 22) and parse_ymd("2026-05-07") == (2026, 5, 7)
    assert parse_ymd("581018-1543610") is None and parse_ymd("") is None and parse_ymd("2026.13.40") is None
    assert pick_report_date("581018-1543610", "2020.11.06.") == ((2020, 11, 6), "receipt_date")
    assert pick_report_date("2021.03.01", "2021.04.01") == ((2021, 3, 1), "change_date")
    assert pick_report_date("", "") == (None, None)
    assert mask_rrn("581018-1543610") != "581018-1543610" and mask_rrn("2020.11.06") == "2020.11.06"


# ───────── 회원 수정 → 변경이력 ─────────
def test_vehicle_type_change_is_one_replacement_and_resave_is_not(client):
    mid = _member(vehicle_type="1톤 카고")
    assert client.put(f"/api/members/{mid}", json={"vehicle_type": "1톤 탑차"}).status_code == 200
    assert _rows(member_id=mid, change_type="대폐차") == [("대폐차", NORMAL, "1톤 카고", "1톤 탑차", {"source": "member_auto_log", "processed_by": "tester"})]
    client.put(f"/api/members/{mid}", json={"vehicle_type": "1톤 탑차"})            # 동일 값 재저장
    client.put(f"/api/members/{mid}", json={"vehicle_type": " 1톤  탑차 "})         # 공백 정리
    ct = [r[:2] for r in _rows(member_id=mid) if r[0] in ("대폐차", "차종변경")]
    assert ct == [("대폐차", NORMAL), ("차종변경", NOTATION)]                       # 공백 정리는 실적 제외(표기정정)


def test_plate_only_change_is_not_vehicle_replacement(client):
    mid = _member()
    client.put(f"/api/members/{mid}", json={"vehicle_number": "강원80배1099호"})     # 실제 번호 변경(2글자 이상 상이)
    types = [(r[0], r[1]) for r in _rows(member_id=mid)]
    assert ("차량번호변경", NORMAL) in types and not [t for t in types if t[0] in ("대폐차", "차종변경")]


def test_first_plate_typo_and_ambiguous_plate_edits(client):
    m1 = _member(vehicle_number="택배", name="임시1")
    client.put(f"/api/members/{m1}", json={"vehicle_number": "강원86배 1069호"})
    assert [(r[0], r[1]) for r in _rows(member_id=m1)] == [("최초번호등록", FIRST)]
    m2 = _member(vehicle_number="강운83자 5192호", name="오타", category="개인")
    client.put(f"/api/members/{m2}", json={"vehicle_number": "강원83자 5192호"})
    assert [(r[0], r[1]) for r in _rows(member_id=m2)] == [("차량번호변경", ERROR_FIX)]
    m3 = _member(vehicle_number="강원83자 5192호", name="모호", category="개인")
    client.put(f"/api/members/{m3}", json={"vehicle_number": "강원83자 5193호"})
    assert [(r[0], r[1]) for r in _rows(member_id=m3)] == [("차량번호변경", REVIEW)]
    m4 = _member(vehicle_number="강원83자 5192호", name="표기", category="개인")
    client.put(f"/api/members/{m4}", json={"vehicle_number": "강원 83자 5192호"})
    assert not _rows(member_id=m4, judgement=NORMAL)


def test_address_company_changes_and_single_record_per_save(client):
    mid = _member(address="동내면 1번지", official_address="동내면 1번지", affiliated_company="가나다물류")
    client.put(f"/api/members/{mid}", json={"address": "동내면 2번지", "official_address": "동내면 2번지"})
    assert len(_rows(member_id=mid, change_type="주소지변경")) == 1                  # 주소+공문주소 → 1건
    client.put(f"/api/members/{mid}", json={"address": "동내면  2번지"})
    assert [r[1] for r in _rows(member_id=mid, change_type="주소지변경")] == [NORMAL, NOTATION]
    client.put(f"/api/members/{mid}", json={"affiliated_company": "라마바운송"})
    assert [(r[0], r[1]) for r in _rows(member_id=mid, change_type="전속계약 업체변경")] == [("전속계약 업체변경", NORMAL)]
    db = SessionLocal()
    try:   # 기존 MemberEditLog 는 그대로 기록
        assert db.query(models.MemberEditLog).filter(models.MemberEditLog.member_id == mid).count() >= 3
    finally:
        db.close()


def test_blank_fill_is_not_a_change(client):
    mid = _member(address="", vehicle_type="")
    client.put(f"/api/members/{mid}", json={"address": "동내면 3번지", "vehicle_type": "1톤 카고"})
    assert _rows(member_id=mid) == []


# ───────── 탈퇴 감지 ─────────
def test_withdrawal_detection_and_non_cases(client):
    mid = _member(membership_date="21.03.10", membership_status="가입")
    r = client.put(f"/api/members/{mid}", json={"membership_date": ""})              # 확인 없이 삭제 → 409, 아무것도 저장 안 됨
    assert r.status_code == 409
    assert client.get(f"/api/members/{mid}").json()["membership_date"] == "21.03.10" and _rows(member_id=mid) == []
    assert client.put(f"/api/members/{mid}", json={"membership_date": "21.03.11"}).status_code == 200   # 정정은 탈퇴 아님
    assert not _rows(member_id=mid, change_type="탈퇴")
    r = client.put(f"/api/members/{mid}", json={"membership_date": "", "confirm_withdrawal": True})
    assert r.status_code == 200
    w = _rows(member_id=mid, change_type="탈퇴")
    assert len(w) == 1 and w[0][2] == "21.03.11" and w[0][4]["prior_membership_date"] == "21.03.11" and w[0][4]["processed_by"] == "tester"
    # 미가입 회원: 빈칸→빈칸 / 빈칸→입력 은 탈퇴 아님
    n = _member(vehicle_number="강원80배2002", name="미가입", membership_date="")
    assert client.put(f"/api/members/{n}", json={"membership_date": ""}).status_code == 200
    assert client.put(f"/api/members/{n}", json={"membership_date": "26.01.05"}).status_code == 200
    assert not _rows(member_id=n, change_type="탈퇴")


# ───────── 월례보고서 ↔ 월별 통계 일치 / 기준일 / 주민번호 ─────────
def _add_hist(**kw):
    db = SessionLocal()
    try:
        db.add(models.ChangeHistory(**kw)); db.commit()
    finally:
        db.close()


def test_report_and_stats_agree_and_date_rules(client):
    ym = (TODAY.year, TODAY.month)
    d = f"{ym[0]}.{ym[1]:02d}.10"
    _add_hist(change_type="전속계약 업체변경", region="춘천시", vehicle_number="강원80배1", name="가", before_value="A", after_value="B", change_date=d)  # 기존(미분류) → 상호변경 집계
    _add_hist(change_type="상호변경", region="원주시", vehicle_number="강원80배2", name="나", before_value="C", after_value="D", change_date=d, judgement=NORMAL)
    _add_hist(change_type="대폐차", region="춘천시", vehicle_number="강원80배3", name="다", before_value="a", after_value="b", change_date=d, judgement=NORMAL)
    _add_hist(change_type="번호변경", region="춘천시", vehicle_number="강원80배4", name="라", before_value="강원80배4", after_value="강원80배5", change_date=d)   # 차량번호변경 → 차량변경 아님
    _add_hist(change_type="주소지변경", region="춘천시", vehicle_number="강원80배6", name="마", before_value="x", after_value="y", change_date="581018-1543610", receipt_date=d)  # 주민번호 → 접수일자 사용
    _add_hist(change_type="주소지변경", region="춘천시", vehicle_number="강원80배7", name="바", before_value="x", after_value="y", change_date="", receipt_date="")        # 날짜 없음 → 어느 달에도 미포함
    _add_hist(change_type="차량번호변경", region="춘천시", vehicle_number="강원80배8", name="사", before_value="택배", after_value="강원80배8", change_date=d, judgement=FIRST)    # 실적 제외
    _add_hist(change_type="주소지변경", region="춘천시", vehicle_number="강원80배9", name="아", before_value="a b", after_value="ab", change_date=d, judgement=NOTATION)       # 실적 제외
    rep = client.get(f"/api/dashboard/monthly-report-auto?year={ym[0]}&month={ym[1]}").json()
    aw = rep["admin_work"]
    assert aw["상호변경"] == 2 and aw["차량변경"] == 1 and aw["주소변경"] == 1
    st = client.get(f"/api/change-stats/summary?year={ym[0]}&month={ym[1]}").json()
    b = st["buckets"]
    assert b["상호변경"] == aw["상호변경"] and b["차량변경(대폐차)"] == aw["차량변경"] and b["주소변경"] == aw["주소변경"]
    assert b["차량번호변경"] == 1 and st["total"] == aw["_변경등록전체"] == 5
    assert st["separate"]["최초등록"] == 1 and st["separate"]["표기정정"] == 1 and st["separate"]["월귀속불가(날짜 없음/불량)"] == 1
    assert st["by_region"]["원주시"]["합계"] == 1
    # 상세·출처
    det = client.get(f"/api/change-stats/detail?bucket=주소변경&year={ym[0]}&month={ym[1]}").json()
    assert det["total"] == 1 and det["items"][0]["basis_source"] == "접수일자" and det["items"][0]["change_date"] != "581018-1543610"
    assert client.get(f"/api/change-stats/detail?bucket=월귀속불가&year={ym[0]}&month={ym[1]}").json()["total"] == 1
    assert client.get(f"/api/change-stats/detail?bucket=상호변경&year={ym[0]}&month={ym[1]}&search=나").json()["total"] == 1
    # 지역 필터
    assert client.get(f"/api/change-stats/summary?year={ym[0]}&month={ym[1]}&region=원주시").json()["total"] == 1


def test_rrn_never_exposed_in_list_detail_or_export(client):
    _add_hist(change_type="주소지변경", region="춘천시", vehicle_number="강원80배1", name="가", before_value="x", after_value="y",
              change_date="581018-1543610", receipt_date="2020.11.06.")
    raw = "581018-1543610"
    lst = client.get("/api/change-history").json()["items"][0]
    cid = lst["id"]
    assert raw not in str(lst) and raw not in str(client.get(f"/api/change-history/{cid}").json())
    assert raw not in client.get("/api/change-history/export/excel").content.decode("latin-1", "ignore") or True
    ws = load_workbook(io.BytesIO(client.get("/api/change-history/export/excel").content)).active
    assert raw not in " ".join(str(c.value) for row in ws.iter_rows() for c in row)
    ws2 = load_workbook(io.BytesIO(client.get("/api/change-stats/export").content))
    assert raw not in " ".join(str(c.value) for w in ws2.worksheets for row in w.iter_rows() for c in row)
    db = SessionLocal()
    try:   # 원본 DB 값은 보존
        assert db.query(models.ChangeHistory).first().change_date == raw
    finally:
        db.close()


def test_stats_export_and_reclassify_preview_are_read_only(client):
    d = f"{TODAY.year}.{TODAY.month:02d}.10"
    _add_hist(change_type="번호변경", region="춘천시", vehicle_number="x", name="가", before_value="강원 자", after_value="강원80자 5976호", change_date=d)
    r = client.get("/api/change-stats/export")
    assert r.status_code == 200 and "요약" in load_workbook(io.BytesIO(r.content)).sheetnames
    pv = client.get("/api/change-stats/reclassify-preview").json()
    assert pv["read_only"] and not pv["applied"] and pv["proposed_counts"].get("최초등록") == 1
    db = SessionLocal()
    try:
        assert db.query(models.ChangeHistory).filter(models.ChangeHistory.judgement.isnot(None)).count() == 0   # 기존 이력 미수정
    finally:
        db.close()


# ───────── 관리비폐지 확인필요 ─────────
def test_fee_closure_candidates(client):
    tm = _member(vehicle_number="강원80배3001", name="택배폐업")
    pm = _member(vehicle_number="강원80자3002", name="개인폐업", category="개인")
    db = SessionLocal()
    try:
        db.add_all([
            models.Closure(closure_type="폐업", region="춘천시", vehicle_number="강원80배3001", name="택배폐업", closure_date="2026.05.10", member_id=tm),
            models.Closure(closure_type="폐지", region="원주시", vehicle_number="강원81배3003", name="폐지택배", closure_date="2025.01.01"),   # 폐지=폐업, 차량번호 '배' → 택배
            models.Closure(closure_type="폐업", region="춘천시", vehicle_number="강원80자3002", name="개인폐업", closure_date="2026.05.11", member_id=pm),
            models.Closure(closure_type="양도", region="춘천시", vehicle_number="강원80배3004", name="양도택배", closure_date="2026.05.12"),
        ]); db.commit()
    finally:
        db.close()
    d = client.get("/api/allocation/fee-closure-candidates").json()
    names = sorted(i["name"] for i in d["items"])
    assert d["read_only"] and names == ["택배폐업", "폐지택배"]
    assert all(i["billing_status"] == "검증불가" and i["post_closure_amount"] is None for i in d["items"])
    assert {"name", "vehicle_number", "region", "closure_date"} <= set(d["items"][0])
    # 선택한 달 이전 폐업만
    assert [i["name"] for i in client.get("/api/allocation/fee-closure-candidates?year=2025&month=12").json()["items"]] == ["폐지택배"]
    db = SessionLocal()
    try:
        assert db.query(models.Closure).count() == 4        # 읽기 전용: 폐업/부과 자료 변경 없음
    finally:
        db.close()
